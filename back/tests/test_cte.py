import io
import asyncio
import sys
import unittest
from pathlib import Path

import openpyxl
BACK_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACK_DIR))

import main  # noqa: E402
from document_segmentation import _page_classification  # noqa: E402


class CteRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _, cls.services = main.load_fixed_cadastros()

    def test_dacte_is_classified_as_cte(self):
        text = """
        DACTE
        DOCUMENTO AUXILIAR DO CONHECIMENTO DE TRANSPORTE ELETRÔNICO
        CT-E Nº 123456 SÉRIE 1
        """
        self.assertEqual(main.classify(text), "CTE")
        self.assertEqual(_page_classification(text)[0], "CTE")

    def test_cte_row_uses_fixed_totvs_transport_code(self):
        row = main.build_cte_row(
            ["DACTE", "CT-E Nº 123456 SÉRIE 1"],
            self.services,
            "PCM-01",
            "cte.pdf",
            "doc-1",
            91.0,
        )
        self.assertEqual(row["Número da Nota"], "123456")
        self.assertEqual(row["Código no Banco de Dados"], "780007")
        self.assertEqual(row["Descrição no Banco de Dados"], "SERVICO TRANSPORTE")
        self.assertEqual(row["_document_id"], "doc-1")

    def test_review_payload_keeps_cte_out_of_outros(self):
        row = main.build_cte_row(
            ["DACTE", "CT-E Nº 42"], self.services, "PCM", "cte.pdf", "doc-1", 90.0
        )
        session = {
            "modo": "normal",
            "documents": {
                "doc-1": {
                    "filename": "cte.pdf",
                    "pages": [1],
                    "tipo": "CTE",
                    "source_filename": "cte.pdf",
                }
            },
            "nfe_rows": [],
            "nfse_rows": [],
            "cte_rows": [row],
            "outros_rows": [],
            "orcamento_rows": [],
        }
        payload = main._build_review_payload("session-1", session)
        self.assertEqual(payload["resumo"]["cte"], 1)
        self.assertEqual(payload["resumo"]["outros"], 0)
        self.assertEqual(payload["cte_rows"][0]["Código no Banco de Dados"], "780007")

    def test_excel_has_dedicated_cte_sheet_without_confidence_metadata(self):
        session_id = "test-export-cte"
        main.REVISION_SESSIONS[session_id] = {
            "created_at": main.time.time(),
            "documents": {},
        }
        request = main.ExportRequest(**{
                "session_id": session_id,
                "centro_custo": "PCM",
                "nfe_rows": [{
                    "Código do Produto do Emitente": "ABC-001/XP",
                    "Código no Banco de Dados": "TOTVS-900",
                    "Arquivo": "nfe.pdf",
                    "Observações": "Não exportar.",
                }],
                "nfse_rows": [],
                "cte_rows": [{
                    "Código no Banco de Dados": "780007",
                    "Descrição no Banco de Dados": "SERVICO TRANSPORTE",
                    "Confiança": "100%",
                    "confiancas": {"Código no Banco de Dados": 100},
                    "Arquivo": "cte.pdf",
                    "Observações": "Visível apenas durante a revisão.",
                }],
                "outros_rows": [],
                "orcamento_rows": [],
            })
        response = asyncio.run(main.exportar(request))

        async def collect_body():
            return b"".join([chunk async for chunk in response.body_iterator])

        workbook = openpyxl.load_workbook(io.BytesIO(asyncio.run(collect_body())))
        self.assertIn("Transportes (CT-e)", workbook.sheetnames)
        sheet = workbook["Transportes (CT-e)"]
        headers = [cell.value for cell in sheet[1]]
        self.assertNotIn("Confiança", headers)
        self.assertNotIn("confiancas", headers)
        self.assertNotIn("Arquivo", headers)
        self.assertNotIn("Observações", headers)
        self.assertEqual(sheet.cell(2, headers.index("Código no Banco de Dados") + 1).value, "780007")
        product_sheet = workbook["Produtos (NF-e)"]
        product_headers = [cell.value for cell in product_sheet[1]]
        self.assertIn("Código do Produto do Emitente", product_headers)
        self.assertNotIn("Arquivo", product_headers)
        self.assertNotIn("Observações", product_headers)
        self.assertEqual(
            product_sheet.cell(2, product_headers.index("Código do Produto do Emitente") + 1).value,
            "ABC-001/XP",
        )


if __name__ == "__main__":
    unittest.main()
